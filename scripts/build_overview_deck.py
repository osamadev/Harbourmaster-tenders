"""Generate a PowerPoint deck explaining Harbourmaster's architecture and key features."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt, Emu


# ---------- palette ----------
NAVY = RGBColor(0x0B, 0x1F, 0x3A)
TEAL = RGBColor(0x14, 0xB8, 0xA6)
SLATE = RGBColor(0x33, 0x41, 0x55)
LIGHT = RGBColor(0xF1, 0xF5, 0xF9)
GREY = RGBColor(0x64, 0x74, 0x8B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
AMBER = RGBColor(0xF5, 0x9E, 0x0B)
ROSE = RGBColor(0xE1, 0x1D, 0x48)
GREEN = RGBColor(0x16, 0xA3, 0x4A)
INDIGO = RGBColor(0x4F, 0x46, 0xE5)


SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def new_deck() -> Presentation:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs


def add_blank(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[6])


def add_rect(slide, left, top, width, height, fill, line=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(0.75)
    shp.shadow.inherit = False
    return shp


def add_rounded(slide, left, top, width, height, fill, line=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shp.adjustments[0] = 0.12
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(0.75)
    shp.shadow.inherit = False
    return shp


def add_text(slide, left, top, width, height, text, *, size=14, bold=False,
             color=SLATE, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font="Calibri"):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    lines = text.split("\n") if isinstance(text, str) else text
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return tb


def add_bullets(slide, left, top, width, height, items, *, size=14, color=SLATE,
                bold_first_token=False, bullet_color=TEAL, line_spacing=1.15):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.line_spacing = line_spacing
        # bullet glyph
        dot = p.add_run()
        dot.text = "■  "
        dot.font.size = Pt(size)
        dot.font.color.rgb = bullet_color
        dot.font.bold = True
        # body
        if bold_first_token and " — " in item:
            head, tail = item.split(" — ", 1)
            r1 = p.add_run()
            r1.text = head + " — "
            r1.font.name = "Calibri"
            r1.font.size = Pt(size)
            r1.font.bold = True
            r1.font.color.rgb = NAVY
            r2 = p.add_run()
            r2.text = tail
            r2.font.name = "Calibri"
            r2.font.size = Pt(size)
            r2.font.color.rgb = color
        else:
            r = p.add_run()
            r.text = item
            r.font.name = "Calibri"
            r.font.size = Pt(size)
            r.font.color.rgb = color
    return tb


def add_title_band(slide, title, subtitle=None):
    add_rect(slide, 0, 0, SLIDE_W, Inches(0.95), NAVY)
    add_rect(slide, 0, Inches(0.95), SLIDE_W, Inches(0.06), TEAL)
    add_text(slide, Inches(0.5), Inches(0.18), Inches(12.3), Inches(0.55),
             title, size=26, bold=True, color=WHITE,
             anchor=MSO_ANCHOR.MIDDLE)
    if subtitle:
        add_text(slide, Inches(0.5), Inches(1.1), Inches(12.3), Inches(0.4),
                 subtitle, size=14, color=GREY)


def add_footer(slide, page_num):
    add_text(slide, Inches(0.5), Inches(7.15), Inches(6), Inches(0.3),
             "Harbourmaster  |  Workflow governance control plane",
             size=10, color=GREY)
    add_text(slide, Inches(12.0), Inches(7.15), Inches(1.0), Inches(0.3),
             str(page_num), size=10, color=GREY, align=PP_ALIGN.RIGHT)


def add_arrow(slide, x1, y1, x2, y2, color=TEAL, weight=Pt(2.25)):
    line = slide.shapes.add_connector(1, x1, y1, x2, y2)  # straight
    line.line.color.rgb = color
    line.line.width = weight
    # arrow head
    from pptx.oxml.ns import qn
    from lxml import etree
    ln = line.line._get_or_add_ln()
    tail = etree.SubElement(ln, qn("a:tailEnd"))
    tail.set("type", "triangle")
    tail.set("w", "med")
    tail.set("h", "med")
    return line


# =========================================================
# Slides
# =========================================================

def slide_title(prs):
    s = add_blank(prs)
    add_rect(s, 0, 0, SLIDE_W, SLIDE_H, NAVY)
    # decorative band
    add_rect(s, 0, Inches(3.1), SLIDE_W, Inches(0.05), TEAL)
    add_text(s, Inches(0.8), Inches(2.0), Inches(11.7), Inches(0.6),
             "HARBOURMASTER", size=18, bold=True, color=TEAL,
             align=PP_ALIGN.LEFT)
    add_text(s, Inches(0.8), Inches(2.4), Inches(11.7), Inches(0.8),
             "A Workflow Governance Control Plane", size=40, bold=True, color=WHITE)
    add_text(s, Inches(0.8), Inches(3.3), Inches(11.7), Inches(0.6),
             "for Agentic Tender & Contract Review",
             size=24, color=LIGHT)
    add_text(s, Inches(0.8), Inches(4.3), Inches(11.7), Inches(0.4),
             "Built on Google Gemini  ·  LangGraph  ·  Arize Phoenix  ·  Elasticsearch  ·  Streamlit",
             size=14, color=GREY)
    # tagline
    add_rounded(s, Inches(0.8), Inches(5.4), Inches(11.7), Inches(1.0), SLATE)
    add_text(s, Inches(1.0), Inches(5.55), Inches(11.3), Inches(0.7),
             "Two-layer governance for enterprise LLM workflows:\n"
             "an inline guard evaluator and a workflow-layer governor.",
             size=16, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)


def slide_problem(prs):
    s = add_blank(prs)
    add_title_band(s, "The Problem", "Why governing agentic LLM workflows is hard")

    items = [
        "Prompt injection — uploaded tenders can carry adversarial instructions that hijack the agent.",
        "Credential & key leakage — bearer tokens, secrets and API keys flow through chat traffic.",
        "Policy blindness — generic LLMs do not know your corporate procurement and compliance rules.",
        "No audit trail — opaque multi-agent calls are impossible for risk and legal to verify.",
        "All-or-nothing autonomy — workflows either run fully unattended or require human approval on every step.",
    ]
    add_bullets(s, Inches(0.7), Inches(1.6), Inches(7.6), Inches(5.0),
                items, size=18, bold_first_token=False, line_spacing=1.35)

    # right pillar callout
    add_rounded(s, Inches(8.8), Inches(1.6), Inches(4.0), Inches(5.0), LIGHT)
    add_text(s, Inches(9.0), Inches(1.8), Inches(3.6), Inches(0.5),
             "What enterprises need", size=16, bold=True, color=NAVY)
    add_bullets(s, Inches(9.0), Inches(2.35), Inches(3.6), Inches(4.3),
                ["Conversation-level safety",
                 "Workflow-level risk control",
                 "Policy-aware reasoning",
                 "Human review on demand",
                 "Full audit log"],
                size=14, bullet_color=AMBER, line_spacing=1.4)

    add_footer(s, 2)


def slide_solution(prs):
    s = add_blank(prs)
    add_title_band(s, "The Solution",
                   "Governance at two layers, around every single LLM call")

    # two cards
    card_w = Inches(5.8)
    card_h = Inches(4.6)
    top = Inches(1.7)

    # Inline Guard
    add_rounded(s, Inches(0.6), top, card_w, card_h, LIGHT)
    add_rect(s, Inches(0.6), top, card_w, Inches(0.55), NAVY)
    add_text(s, Inches(0.8), top + Inches(0.07), card_w - Inches(0.4), Inches(0.4),
             "Inline Guard  ·  LLM-as-Judge",
             size=16, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    add_bullets(s, Inches(0.85), top + Inches(0.75), card_w - Inches(0.5), Inches(3.7),
                [
                    "Evaluates every prompt and model response",
                    "Verdicts: ALLOW · HUMAN_REVIEW · DENY",
                    "Detects injection, exfiltration, credentials",
                    "Runs before workflow and on completions",
                    "Phoenix guard spans (direction=guard)",
                ],
                size=14, bullet_color=ROSE, line_spacing=1.45)

    # Harbourmaster
    add_rounded(s, Inches(6.9), top, card_w, card_h, LIGHT)
    add_rect(s, Inches(6.9), top, card_w, Inches(0.55), TEAL)
    add_text(s, Inches(7.1), top + Inches(0.07), card_w - Inches(0.4), Inches(0.4),
             "Harbourmaster  ·  Workflow Governor",
             size=16, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    add_bullets(s, Inches(7.15), top + Inches(0.75), card_w - Inches(0.5), Inches(3.7),
                [
                    "Multi-agent contract review on LangGraph",
                    "Specialist analysts + critic/verifier loop",
                    "Risk-based routing with human review",
                    "Counter-clause negotiator for high-risk text",
                    "Corporate policy library with scope tags",
                ],
                size=14, bullet_color=TEAL, line_spacing=1.45)

    # bottom note
    add_rounded(s, Inches(0.6), Inches(6.5), Inches(12.1), Inches(0.55), SLATE)
    add_text(s, Inches(0.8), Inches(6.52), Inches(11.7), Inches(0.5),
             "A risky indemnity clause is not necessarily an attack — the inline guard may allow it,"
             " while Harbourmaster still pauses the workflow because business risk is high.",
             size=12, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)

    add_footer(s, 3)


def slide_two_layers_table(prs):
    s = add_blank(prs)
    add_title_band(s, "Two Governance Layers",
                   "Different decisions, different mechanisms, complementary scope")

    headers = ["Area", "Inline Guard", "Harbourmaster"]
    rows = [
        ["Layer", "Safety evaluator", "Workflow governor"],
        ["Decides", "Is this prompt/response an attack?",
         "Does this business workflow need review?"],
        ["Catches", "Injection, exfiltration, credentials, dangerous commands",
         "Contract risk, compliance gaps, policy conflicts"],
        ["Mechanism", "LLM-as-judge, Phoenix guard spans",
         "LangGraph routing, verifier loop, human interrupt"],
        ["Example", "Poisoned tender text is denied",
         "Unlimited liability clause is escalated"],
    ]

    left = Inches(0.6)
    top = Inches(1.6)
    total_w = Inches(12.1)
    col_w = [Inches(2.2), Inches(4.95), Inches(4.95)]
    header_h = Inches(0.55)
    row_h = Inches(0.85)

    # header
    x = left
    for i, h in enumerate(headers):
        bg = NAVY if i == 0 else (ROSE if i == 1 else TEAL)
        add_rect(s, x, top, col_w[i], header_h, bg)
        add_text(s, x + Inches(0.15), top, col_w[i] - Inches(0.2), header_h,
                 h, size=14, bold=True, color=WHITE,
                 anchor=MSO_ANCHOR.MIDDLE)
        x += col_w[i]

    # rows
    y = top + header_h
    for r, row in enumerate(rows):
        x = left
        bg = WHITE if r % 2 == 0 else LIGHT
        for i, cell in enumerate(row):
            add_rect(s, x, y, col_w[i], row_h, bg, line=GREY)
            bold = i == 0
            col = NAVY if i == 0 else SLATE
            add_text(s, x + Inches(0.15), y, col_w[i] - Inches(0.2), row_h,
                     cell, size=12, bold=bold, color=col,
                     anchor=MSO_ANCHOR.MIDDLE)
            x += col_w[i]
        y += row_h

    add_footer(s, 4)


def slide_architecture(prs):
    s = add_blank(prs)
    add_title_band(s, "High-Level Architecture",
                   "Inline guard, direct Gemini calls, Phoenix telemetry, Elastic memory")

    # nodes
    def box(left, top, w, h, label, fill, sub=None, text_color=WHITE):
        add_rounded(s, left, top, w, h, fill)
        add_text(s, left, top + (Inches(0.05) if sub else Inches(0)),
                 w, Inches(0.45),
                 label, size=14, bold=True, color=text_color,
                 align=PP_ALIGN.CENTER,
                 anchor=MSO_ANCHOR.MIDDLE)
        if sub:
            add_text(s, left, top + Inches(0.42), w, Inches(0.3),
                     sub, size=10, color=LIGHT,
                     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    # column 1 — inputs
    box(Inches(0.5), Inches(1.7), Inches(2.4), Inches(0.75),
        "Tender PDF / MD / TXT", SLATE)
    box(Inches(0.5), Inches(2.7), Inches(2.4), Inches(0.75),
        "Corporate Policies", SLATE,
        sub="configs/corporate_policies/*.json")

    # column 2 — UI / workflow
    box(Inches(3.5), Inches(1.7), Inches(3.0), Inches(0.75),
        "Streamlit UI", INDIGO,
        sub="reviewer · policies · dashboard")
    box(Inches(3.5), Inches(2.85), Inches(3.0), Inches(2.7),
        "LangGraph Workflow", TEAL)
    add_bullets(s, Inches(3.6), Inches(3.3), Inches(2.8), Inches(2.2),
                ["segment clauses",
                 "5 specialist analysts",
                 "aggregate + verifier loop",
                 "risk routing",
                 "negotiator + drafter"],
                size=11, color=WHITE, bullet_color=LIGHT, line_spacing=1.25)

    # column 3 — governed call path
    box(Inches(7.0), Inches(1.7), Inches(2.9), Inches(0.75),
        "Inline Guard", ROSE, sub="guard.py")
    box(Inches(7.0), Inches(2.7), Inches(2.9), Inches(0.75),
        "GovernedLLM", SLATE, sub="direct Gemini client")
    box(Inches(7.0), Inches(3.7), Inches(2.9), Inches(0.75),
        "Gemini API", NAVY,
        sub="gemini-2.5-pro / flash")

    # column 4 — observability + memory
    box(Inches(10.2), Inches(1.7), Inches(2.7), Inches(0.75),
        "Arize Phoenix", INDIGO,
        sub=":6006  OTLP traces")
    box(Inches(10.2), Inches(2.7), Inches(2.7), Inches(0.75),
        "Elasticsearch", TEAL, sub=":9200  precedent search")
    box(Inches(10.2), Inches(3.7), Inches(2.7), Inches(0.75),
        "Governance Dashboard", INDIGO)
    box(Inches(10.2), Inches(4.7), Inches(2.7), Inches(0.75),
        "Human Review", GREEN, sub="LangGraph interrupt()")

    # arrows
    add_arrow(s, Inches(2.9), Inches(2.07), Inches(3.5), Inches(2.07))
    add_arrow(s, Inches(2.9), Inches(3.07), Inches(3.5), Inches(3.5))
    add_arrow(s, Inches(6.5), Inches(2.07), Inches(7.0), Inches(2.07))
    add_arrow(s, Inches(6.5), Inches(4.2), Inches(7.0), Inches(2.3),
              color=GREY)
    add_arrow(s, Inches(8.45), Inches(2.45), Inches(8.45), Inches(2.7), color=ROSE)
    add_arrow(s, Inches(8.45), Inches(3.45), Inches(8.45), Inches(3.7), color=SLATE)
    add_arrow(s, Inches(9.9), Inches(2.07), Inches(10.2), Inches(2.07), color=INDIGO)
    add_arrow(s, Inches(9.9), Inches(3.2), Inches(10.2), Inches(2.9), color=TEAL)

    # caption
    add_rounded(s, Inches(0.5), Inches(6.3), Inches(12.4), Inches(0.7), LIGHT)
    add_text(s, Inches(0.7), Inches(6.35), Inches(12.0), Inches(0.6),
             "Request path:  guard  →  GovernedLLM  →  Gemini API"
             "       ·       OpenInference spans → Phoenix  ·  policies & reviews → Elastic.",
             size=12, color=NAVY, anchor=MSO_ANCHOR.MIDDLE)

    add_footer(s, 5)


def slide_langgraph_flow(prs):
    s = add_blank(prs)
    add_title_band(s, "LangGraph Multi-Agent Workflow",
                   "Deterministic routing with a critic loop and a human-in-the-loop branch")

    def node(left, top, w, h, label, fill, text_color=WHITE, size=12):
        add_rounded(s, left, top, w, h, fill)
        add_text(s, left, top, w, h, label, size=size, bold=True,
                 color=text_color, align=PP_ALIGN.CENTER,
                 anchor=MSO_ANCHOR.MIDDLE)

    y0 = Inches(1.7)
    node(Inches(0.5), y0, Inches(1.7), Inches(0.7), "START", GREEN)
    node(Inches(2.4), y0, Inches(1.9), Inches(0.7), "segment", SLATE)
    node(Inches(4.5), y0, Inches(2.2), Inches(0.7), "specialists (×5)", TEAL)
    node(Inches(6.9), y0, Inches(1.9), Inches(0.7), "aggregate", SLATE)
    node(Inches(9.0), y0, Inches(1.9), Inches(0.7), "verifier", AMBER)
    node(Inches(11.1), y0, Inches(1.8), Inches(0.7), "governance", NAVY)

    # revision loop
    node(Inches(4.5), Inches(3.0), Inches(2.2), Inches(0.7),
         "revision (++round)", INDIGO)

    # governance branches
    node(Inches(9.4), Inches(3.0), Inches(2.5), Inches(0.7),
         "human_review  (interrupt)", ROSE)
    node(Inches(6.5), Inches(4.3), Inches(2.5), Inches(0.7),
         "negotiator", TEAL)
    node(Inches(9.4), Inches(4.3), Inches(2.5), Inches(0.7),
         "draft", SLATE)
    node(Inches(11.5), Inches(5.6), Inches(1.4), Inches(0.7), "END", GREEN)

    # arrows linear
    add_arrow(s, Inches(2.2), Inches(2.05), Inches(2.4), Inches(2.05))
    add_arrow(s, Inches(4.3), Inches(2.05), Inches(4.5), Inches(2.05))
    add_arrow(s, Inches(6.7), Inches(2.05), Inches(6.9), Inches(2.05))
    add_arrow(s, Inches(8.8), Inches(2.05), Inches(9.0), Inches(2.05))
    add_arrow(s, Inches(10.9), Inches(2.05), Inches(11.1), Inches(2.05))

    # verifier → revision (loop)
    add_arrow(s, Inches(9.6), Inches(2.4), Inches(9.6), Inches(2.7), color=AMBER)
    add_arrow(s, Inches(9.6), Inches(2.7), Inches(5.6), Inches(2.7), color=AMBER)
    add_arrow(s, Inches(5.6), Inches(2.7), Inches(5.6), Inches(3.0), color=AMBER)
    # revision → specialists
    add_arrow(s, Inches(5.6), Inches(3.7), Inches(5.6), Inches(2.4), color=AMBER)
    add_text(s, Inches(7.0), Inches(2.55), Inches(2.4), Inches(0.3),
             "if pending revisions  &  round < max",
             size=10, color=AMBER, align=PP_ALIGN.CENTER)

    # governance → human_review
    add_arrow(s, Inches(12.0), Inches(2.4), Inches(11.5), Inches(3.0), color=ROSE)
    add_text(s, Inches(11.7), Inches(2.65), Inches(2.0), Inches(0.3),
             "high risk / blocked", size=10, color=ROSE)

    # governance → negotiator (low / medium)
    add_arrow(s, Inches(11.8), Inches(2.4), Inches(7.8), Inches(4.3), color=TEAL)
    add_text(s, Inches(7.8), Inches(3.6), Inches(2.3), Inches(0.3),
             "low / medium risk", size=10, color=TEAL)

    # human_review → negotiator
    add_arrow(s, Inches(10.6), Inches(3.7), Inches(8.0), Inches(4.3), color=ROSE)
    # negotiator → draft
    add_arrow(s, Inches(9.0), Inches(4.65), Inches(9.4), Inches(4.65))
    # draft → END
    add_arrow(s, Inches(11.9), Inches(4.65), Inches(12.2), Inches(5.6))

    add_footer(s, 6)


def slide_specialists(prs):
    s = add_blank(prs)
    add_title_band(s, "Specialist Analysts",
                   "Parallel domain reviewers — each producing structured findings")

    cards = [
        ("Legal", "Liability, indemnity, governing law, dispute resolution.", NAVY),
        ("Financial", "Payment terms, penalties, currency, escalation clauses.", TEAL),
        ("Delivery", "Timelines, SLAs, milestones, acceptance criteria.", INDIGO),
        ("IP / Data", "IP ownership, data residency, confidentiality, GDPR.", AMBER),
        ("Compliance", "Maps clauses to active corporate policies.", ROSE),
    ]

    left = Inches(0.5)
    top = Inches(1.7)
    card_w = Inches(2.45)
    card_h = Inches(3.1)
    gap = Inches(0.1)

    for i, (name, desc, color) in enumerate(cards):
        x = left + (card_w + gap) * i
        add_rounded(s, x, top, card_w, card_h, LIGHT)
        add_rect(s, x, top, card_w, Inches(0.55), color)
        add_text(s, x, top, card_w, Inches(0.55), name,
                 size=15, bold=True, color=WHITE,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        add_text(s, x + Inches(0.2), top + Inches(0.7),
                 card_w - Inches(0.4), Inches(2.3),
                 desc, size=12, color=SLATE)
        # score chip
        add_rounded(s, x + Inches(0.2), top + Inches(2.45),
                    card_w - Inches(0.4), Inches(0.5), color)
        add_text(s, x + Inches(0.2), top + Inches(2.45),
                 card_w - Inches(0.4), Inches(0.5),
                 "Weighted risk score", size=11, bold=True, color=WHITE,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    # bottom note
    add_rounded(s, Inches(0.5), Inches(5.1), Inches(12.4), Inches(1.85), LIGHT)
    add_text(s, Inches(0.7), Inches(5.2), Inches(12.0), Inches(0.4),
             "How findings are structured", size=14, bold=True, color=NAVY)
    add_bullets(s, Inches(0.7), Inches(5.55), Inches(12.0), Inches(1.4),
                [
                    "Each finding cites a stable clause_id and a risk_level (low / medium / high).",
                    "Compliance findings additionally cite policy_id and policy_reference.",
                    "Specialists run in parallel via ThreadPoolExecutor for low latency.",
                    "Overall risk is a configurable weighted average across specialists.",
                ],
                size=12, line_spacing=1.25)

    add_footer(s, 7)


def slide_verifier(prs):
    s = add_blank(prs)
    add_title_band(s, "Critic / Verifier Loop",
                   "Self-correction before a finding ever reaches the reviewer")

    # left: actions
    add_text(s, Inches(0.7), Inches(1.7), Inches(5.5), Inches(0.4),
             "Verifier inspects every merged finding", size=16, bold=True, color=NAVY)
    add_bullets(s, Inches(0.7), Inches(2.15), Inches(6.0), Inches(4.5),
                [
                    "keep — finding is correct, propagate it",
                    "revise — return finding to its specialist with a reason; counts as a revision round",
                    "drop — finding is unsupported by clause text, remove it",
                ],
                size=14, line_spacing=1.5, bullet_color=AMBER)

    add_rounded(s, Inches(0.7), Inches(5.6), Inches(6.0), Inches(1.3), LIGHT)
    add_text(s, Inches(0.9), Inches(5.7), Inches(5.6), Inches(0.4),
             "Guardrails", size=13, bold=True, color=NAVY)
    add_bullets(s, Inches(0.9), Inches(6.05), Inches(5.6), Inches(0.85),
                ["Bounded by VERIFIER_MAX_REVISIONS",
                 "Deterministic merge keyed on (specialist, clause_id, risk_level, rationale)"],
                size=11, line_spacing=1.3)

    # right: mini state machine
    nx = Inches(7.0)
    add_rounded(s, nx, Inches(1.7), Inches(5.7), Inches(5.2), LIGHT)
    add_text(s, nx, Inches(1.78), Inches(5.7), Inches(0.4),
             "Verifier decision flow", size=13, bold=True, color=NAVY,
             align=PP_ALIGN.CENTER)

    def vnode(left, top, w, h, label, fill, color=WHITE):
        add_rounded(s, left, top, w, h, fill)
        add_text(s, left, top, w, h, label, size=12, bold=True, color=color,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    vnode(Inches(7.2), Inches(2.3), Inches(2.0), Inches(0.6), "Merged findings", SLATE)
    vnode(Inches(10.4), Inches(2.3), Inches(2.0), Inches(0.6), "Verifier", AMBER)
    add_arrow(s, Inches(9.2), Inches(2.6), Inches(10.4), Inches(2.6))

    vnode(Inches(7.2), Inches(3.4), Inches(1.6), Inches(0.6), "keep", GREEN)
    vnode(Inches(9.1), Inches(3.4), Inches(1.6), Inches(0.6), "revise", INDIGO)
    vnode(Inches(11.0), Inches(3.4), Inches(1.4), Inches(0.6), "drop", ROSE)

    add_arrow(s, Inches(11.0), Inches(2.9), Inches(8.0), Inches(3.4), color=GREEN)
    add_arrow(s, Inches(11.2), Inches(2.9), Inches(9.9), Inches(3.4), color=INDIGO)
    add_arrow(s, Inches(11.4), Inches(2.9), Inches(11.7), Inches(3.4), color=ROSE)

    vnode(Inches(7.2), Inches(4.5), Inches(5.2), Inches(0.6),
          "specialists (re-run only revised targets)", INDIGO)
    add_arrow(s, Inches(9.9), Inches(4.0), Inches(9.9), Inches(4.5), color=INDIGO)
    add_arrow(s, Inches(9.9), Inches(5.1), Inches(11.4), Inches(2.9), color=INDIGO)

    vnode(Inches(7.2), Inches(5.6), Inches(5.2), Inches(0.6),
          "governance routing", NAVY)
    add_arrow(s, Inches(8.0), Inches(4.0), Inches(8.0), Inches(5.6), color=GREEN)

    add_footer(s, 8)


def slide_human_loop(prs):
    s = add_blank(prs)
    add_title_band(s, "Human-in-the-Loop Governance",
                   "Risk-based routing with a true LangGraph interrupt()")

    add_bullets(s, Inches(0.7), Inches(1.7), Inches(6.0), Inches(4.5),
                [
                    "Overall risk ≥ REVIEW_RISK_THRESHOLD pauses the workflow.",
                    "Inline guard DENY verdicts also force human review.",
                    "interrupt() emits full state: clauses, findings, policies, notes.",
                    "Reviewer can approve, reject, or annotate per clause.",
                    "Decision flows downstream into the counter-clause negotiator.",
                    "On reject, the drafter records 'REVIEW REJECTED by …' instead of a summary.",
                ],
                size=14, line_spacing=1.5)

    # right card — interrupt payload
    nx = Inches(7.1)
    add_rounded(s, nx, Inches(1.7), Inches(5.7), Inches(5.2), LIGHT)
    add_text(s, nx + Inches(0.2), Inches(1.85), Inches(5.4), Inches(0.4),
             "Reviewer payload (LangGraph interrupt)",
             size=14, bold=True, color=NAVY)
    add_text(s, nx + Inches(0.2), Inches(2.25), Inches(5.4), Inches(4.5),
             "type: tender_review_required\n"
             "overall_risk\n"
             "blocked  ·  block_reason\n"
             "findings (verified)\n"
             "clauses (segmented)\n"
             "specialist_findings (per analyst)\n"
             "verifier_notes\n"
             "corporate_policies (active)\n"
             "prompt: 'Approve, reject, or annotate.'",
             size=12, color=SLATE, font="Consolas")

    add_footer(s, 9)


def slide_inline_guard(prs):
    s = add_blank(prs)
    add_title_band(s, "Inline Guard — LLM-as-Judge",
                   "Safety evaluation before workflow execution and on every completion")

    add_bullets(s, Inches(0.7), Inches(1.7), Inches(7.5), Inches(5.0),
                [
                    "Blocks prompt injection, role override, system-prompt extraction.",
                    "Detects credential & API-key exfiltration in both directions.",
                    "Verdicts: ALLOW · HUMAN_REVIEW · DENY.",
                    "Dedicated agent: governance-guard-v1 (Gemini Flash, temp=0).",
                    "Emits Phoenix spans with direction=guard for dashboard charts.",
                    "Red-team suite validates guard outcomes + syncs Phoenix experiments.",
                ],
                size=14, line_spacing=1.45)

    # right — verdict flow
    nx = Inches(8.6)
    add_rounded(s, nx, Inches(1.7), Inches(4.2), Inches(5.2), LIGHT)
    add_text(s, nx + Inches(0.2), Inches(1.85), Inches(3.9), Inches(0.4),
             "Verdict → Workflow", size=14, bold=True, color=NAVY)

    add_rounded(s, nx + Inches(0.2), Inches(2.3), Inches(3.8), Inches(1.0), WHITE, line=GREY)
    add_text(s, nx + Inches(0.3), Inches(2.35), Inches(3.6), Inches(0.4),
             "ALLOW", size=13, bold=True, color=GREEN)
    add_text(s, nx + Inches(0.3), Inches(2.65), Inches(3.6), Inches(0.6),
             "Workflow proceeds normally.",
             size=11, color=SLATE)

    add_rounded(s, nx + Inches(0.2), Inches(3.45), Inches(3.8), Inches(1.0), WHITE, line=GREY)
    add_text(s, nx + Inches(0.3), Inches(3.5), Inches(3.6), Inches(0.4),
             "HUMAN_REVIEW", size=13, bold=True, color=AMBER)
    add_text(s, nx + Inches(0.3), Inches(3.8), Inches(3.6), Inches(0.6),
             "Forces LangGraph interrupt().",
             size=11, color=SLATE)

    add_rounded(s, nx + Inches(0.2), Inches(4.6), Inches(3.8), Inches(1.0), WHITE, line=GREY)
    add_text(s, nx + Inches(0.3), Inches(4.65), Inches(3.6), Inches(0.4),
             "DENY", size=13, bold=True, color=ROSE)
    add_text(s, nx + Inches(0.3), Inches(4.95), Inches(3.6), Inches(0.6),
             "Blocks request; recorded as guard denial span.",
             size=11, color=SLATE)

    add_rounded(s, nx + Inches(0.2), Inches(5.75), Inches(3.8), Inches(1.05), SLATE)
    add_text(s, nx + Inches(0.3), Inches(5.8), Inches(3.6), Inches(0.4),
             "Implementation", size=12, bold=True, color=WHITE)
    add_text(s, nx + Inches(0.3), Inches(6.1), Inches(3.6), Inches(0.7),
             "harbourmaster/guard.py",
             size=10, color=LIGHT, font="Consolas")

    add_footer(s, 10)


def slide_phoenix_elastic(prs):
    s = add_blank(prs)
    add_title_band(s, "Phoenix & Elastic",
                   "Trace-level observability and searchable procurement memory")

    # diagram
    y = Inches(1.8)
    def chip(left, top, w, h, label, fill, sub=None, color=WHITE):
        add_rounded(s, left, top, w, h, fill)
        add_text(s, left, top, w, h - (Inches(0.3) if sub else Inches(0)),
                 label, size=14, bold=True, color=color,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        if sub:
            add_text(s, left, top + h - Inches(0.35), w, Inches(0.3),
                     sub, size=10, color=LIGHT,
                     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    chip(Inches(0.7), y, Inches(2.3), Inches(1.1), "Workflow", TEAL,
         sub="LangGraph + GovernedLLM")
    chip(Inches(3.6), y, Inches(2.3), Inches(1.1), "Phoenix", INDIGO,
         sub=":6006  ·  OTLP spans")
    chip(Inches(6.5), y, Inches(2.3), Inches(1.1), "Dashboard", INDIGO,
         sub="Governance UI")
    chip(Inches(9.4), y, Inches(2.3), Inches(1.1), "Elastic", TEAL,
         sub=":9200  ·  search")
    chip(Inches(11.9), y + Inches(0.1), Inches(1.0), Inches(0.9), "✓", GREEN)

    add_arrow(s, Inches(3.0), y + Inches(0.55), Inches(3.6), y + Inches(0.55))
    add_arrow(s, Inches(5.9), y + Inches(0.55), Inches(6.5), y + Inches(0.55))
    add_arrow(s, Inches(8.8), y + Inches(0.55), Inches(9.4), y + Inches(0.55))

    add_text(s, Inches(0.7), Inches(3.2), Inches(12.0), Inches(0.4),
             "What each layer provides", size=16, bold=True, color=NAVY)
    add_bullets(s, Inches(0.7), Inches(3.65), Inches(7.5), Inches(3.0),
                [
                    "Phoenix: OpenInference traces for every model call and guard verdict.",
                    "Dashboard reads /v1/projects/{project}/spans for risk and denial charts.",
                    "Elastic: semantic search over policies, tenders, clauses, red-team cases.",
                    "Governance Copilot queries Phoenix MCP + Elastic MCP.",
                ],
                size=13, line_spacing=1.45)

    # partner stack card
    add_rounded(s, Inches(8.4), Inches(3.55), Inches(4.5), Inches(3.35), LIGHT)
    add_text(s, Inches(8.6), Inches(3.7), Inches(4.1), Inches(0.4),
             "Partner integrations", size=14, bold=True, color=NAVY)
    add_text(s, Inches(8.6), Inches(4.1), Inches(4.1), Inches(3.0),
             "Arize Phoenix — traces, experiments,\nred-team regression sync.\n\n"
             "Elasticsearch — procurement memory,\nprecedent search, review artifacts.",
             size=12, color=SLATE)

    add_footer(s, 11)


def slide_policies(prs):
    s = add_blank(prs)
    add_title_band(s, "Corporate Policy Library",
                   "Persisted, scoped, and surfaced inside the workflow")

    add_bullets(s, Inches(0.7), Inches(1.7), Inches(6.5), Inches(5.0),
                [
                    "Stored as JSON under configs/corporate_policies/.",
                    "Managed via the Streamlit 'Corporate Policies' page.",
                    "Create, edit, delete, activate / deactivate, scope-tag.",
                    "PDF upload extracts policy text via pypdf.",
                    "Active policies are injected into the LangGraph state.",
                    "Compliance findings cite policy_id and policy_reference.",
                ],
                size=14, line_spacing=1.45)

    # right — policy JSON sample
    nx = Inches(7.6)
    add_rounded(s, nx, Inches(1.7), Inches(5.2), Inches(5.2), SLATE)
    add_text(s, nx + Inches(0.2), Inches(1.85), Inches(5.0), Inches(0.4),
             "Policy schema", size=14, bold=True, color=WHITE)
    sample = (
        "{\n"
        '  "id": "starter-procurement-policy",\n'
        '  "title": "Starter Corporate Procurement Policy",\n'
        '  "version": "1.0",\n'
        '  "scope_tags": [\n'
        '     "compliance", "legal", "financial"\n'
        '  ],\n'
        '  "active": true,\n'
        '  "body": "Policy text..."\n'
        "}"
    )
    add_text(s, nx + Inches(0.25), Inches(2.3), Inches(4.9), Inches(4.4),
             sample, size=12, color=LIGHT, font="Consolas")

    add_footer(s, 12)


def slide_observability(prs):
    s = add_blank(prs)
    add_title_band(s, "Observability & Red Team",
                   "Built-in dashboards and an adversarial scorecard")

    # two cards side by side
    left_x = Inches(0.6)
    right_x = Inches(6.95)
    top = Inches(1.7)
    cw = Inches(5.75)
    ch = Inches(5.25)

    # Governance Dashboard
    add_rounded(s, left_x, top, cw, ch, LIGHT)
    add_rect(s, left_x, top, cw, Inches(0.55), INDIGO)
    add_text(s, left_x, top, cw, Inches(0.55), "Governance Dashboard",
             size=15, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_bullets(s, left_x + Inches(0.25), top + Inches(0.75),
                cw - Inches(0.5), ch - Inches(1.0),
                [
                    "Reads Phoenix project spans via REST API.",
                    "Guard categories from direction=guard spans.",
                    "Risk-bearing calls vs system/tool denials.",
                    "Filters by agent, time window, verdict.",
                    "Streamlit page backed by local Phoenix.",
                    "OpenInference spans from telemetry.py.",
                ],
                size=13, line_spacing=1.4, bullet_color=INDIGO)

    # Red Team Scorecard
    add_rounded(s, right_x, top, cw, ch, LIGHT)
    add_rect(s, right_x, top, cw, Inches(0.55), ROSE)
    add_text(s, right_x, top, cw, Inches(0.55), "Red-Team Scorecard",
             size=15, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_bullets(s, right_x + Inches(0.25), top + Inches(0.75),
                cw - Inches(0.5), ch - Inches(1.0),
                [
                    "Runs curated adversarial prompts.",
                    "Measures inline guard ALLOW / DENY outcomes.",
                    "Surfaces per-attack-category pass rates.",
                    "Run via make redteam or the Streamlit page.",
                    "Syncs results to Phoenix experiments.",
                ],
                size=13, line_spacing=1.4, bullet_color=ROSE)

    add_footer(s, 13)


def slide_deployment(prs):
    s = add_blank(prs)
    add_title_band(s, "Containerized Deployment",
                   "Three services: UI, Phoenix, and Elasticsearch")

    # three service cards
    y = Inches(1.7)
    cw = Inches(4.0)
    ch = Inches(2.7)
    gap = Inches(0.1)

    services = [
        ("ui", "Streamlit UI + LangGraph workflow runner", INDIGO,
         ["Port 8501", "Inline guard + GovernedLLM", "GEMINI_API_KEY in .env"]),
        ("phoenix", "Arize Phoenix observability", INDIGO,
         ["Port 6006", "OTLP trace collector", "Governance dashboard source"]),
        ("elastic", "Elasticsearch procurement memory", TEAL,
         ["Port 9200", "Policy & review indexing", "Semantic precedent search"]),
    ]

    for i, (name, desc, color, items) in enumerate(services):
        x = Inches(0.6) + (cw + gap) * i
        add_rounded(s, x, y, cw, ch, LIGHT)
        add_rect(s, x, y, cw, Inches(0.55), color)
        add_text(s, x + Inches(0.2), y, cw - Inches(0.4), Inches(0.55),
                 name, size=15, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
        add_text(s, x + Inches(0.2), y + Inches(0.65),
                 cw - Inches(0.4), Inches(0.4),
                 desc, size=12, color=NAVY, bold=True)
        add_bullets(s, x + Inches(0.2), y + Inches(1.1),
                    cw - Inches(0.4), Inches(1.6), items,
                    size=11, line_spacing=1.3, bullet_color=color)

    # bottom — persistent volumes + quick start
    by = Inches(4.7)
    add_rounded(s, Inches(0.6), by, Inches(6.0), Inches(2.3), LIGHT)
    add_text(s, Inches(0.8), by + Inches(0.1), Inches(5.6), Inches(0.4),
             "Persistent volumes", size=14, bold=True, color=NAVY)
    add_bullets(s, Inches(0.8), by + Inches(0.55), Inches(5.6), Inches(1.7),
                [
                    "review_data — saved contract reviews and history.",
                    "policy_data — corporate policy JSON library.",
                    "Phoenix and Elastic data survive container restarts.",
                ],
                size=12, line_spacing=1.35)

    add_rounded(s, Inches(6.85), by, Inches(6.05), Inches(2.3), SLATE)
    add_text(s, Inches(7.05), by + Inches(0.1), Inches(5.6), Inches(0.4),
             "Quick start", size=14, bold=True, color=WHITE)
    add_text(s, Inches(7.05), by + Inches(0.55), Inches(5.7), Inches(1.7),
             "docker compose --env-file .env.docker up\n"
             "  →  Streamlit UI :8501\n"
             "  →  Phoenix :6006\n"
             "  →  Elasticsearch :9200",
             size=11, color=LIGHT, font="Consolas")

    add_footer(s, 14)


def slide_key_features(prs):
    s = add_blank(prs)
    add_title_band(s, "Key Features — At a Glance",
                   "What you get with Harbourmaster")

    feats = [
        ("Two-layer governance",
         "Inline guard evaluator AND workflow governor in one stack.", ROSE),
        ("Multi-agent review",
         "5 specialist analysts (legal, financial, delivery, IP/data, compliance).", TEAL),
        ("Critic / verifier loop",
         "keep / revise / drop decisions with bounded re-runs.", AMBER),
        ("Human-in-the-loop",
         "True LangGraph interrupt() with full reviewer payload.", INDIGO),
        ("Counter-clause negotiator",
         "Drafts redlines for high-risk clauses automatically.", TEAL),
        ("Corporate policy library",
         "Scoped JSON policies, PDF upload, runtime activation.", NAVY),
        ("Phoenix telemetry",
         "OpenInference traces for every model call and guard verdict.", INDIGO),
        ("Elastic memory",
         "Semantic search over policies, tenders, and review artifacts.", TEAL),
        ("Red-team scorecard",
         "Adversarial prompts measured against the inline guard.", ROSE),
        ("Containerized deploy",
         "ui + phoenix + elastic Compose stack.", NAVY),
    ]

    # 2 columns × 5 rows
    col_w = Inches(6.2)
    row_h = Inches(1.05)
    top = Inches(1.7)
    for i, (name, desc, color) in enumerate(feats):
        col = i % 2
        row = i // 2
        x = Inches(0.5) + col * (col_w + Inches(0.1))
        y = top + row * (row_h + Inches(0.05))
        add_rounded(s, x, y, col_w, row_h, LIGHT)
        add_rect(s, x, y, Inches(0.18), row_h, color)
        add_text(s, x + Inches(0.3), y + Inches(0.07),
                 col_w - Inches(0.4), Inches(0.4),
                 name, size=13, bold=True, color=NAVY)
        add_text(s, x + Inches(0.3), y + Inches(0.45),
                 col_w - Inches(0.4), Inches(0.55),
                 desc, size=11, color=SLATE)

    add_footer(s, 15)


def slide_stack(prs):
    s = add_blank(prs)
    add_title_band(s, "Tech Stack",
                   "Pragmatic, all open-source apart from the Gemini API")

    stack = [
        ("Models", "Google Gemini 2.5 Pro (analysts) · Gemini 2.5 Flash (drafter + guard)", TEAL),
        ("Orchestration", "LangGraph state machine with MemorySaver checkpointer", INDIGO),
        ("Guard", "Inline LLM-as-judge evaluator (guard.py)", ROSE),
        ("Observability", "Arize Phoenix — OpenInference OTLP traces", INDIGO),
        ("Memory", "Elasticsearch — policy, tender, and review indexing", TEAL),
        ("UI", "Streamlit multi-page app (reviewer, dashboard, policies, copilot)", INDIGO),
        ("Ingestion", "pypdf for tender & policy PDF extraction", SLATE),
        ("Runtime", "Python 3.10+ · Docker Compose (ui + phoenix + elastic)", NAVY),
        ("Licence", "MIT (Harbourmaster)", GREEN),
    ]

    top = Inches(1.7)
    row_h = Inches(0.6)
    for i, (label, desc, color) in enumerate(stack):
        y = top + i * (row_h + Inches(0.08))
        add_rounded(s, Inches(0.7), y, Inches(2.6), row_h, color)
        add_text(s, Inches(0.7), y, Inches(2.6), row_h,
                 label, size=13, bold=True, color=WHITE,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        add_rounded(s, Inches(3.4), y, Inches(9.3), row_h, LIGHT)
        add_text(s, Inches(3.6), y, Inches(9.0), row_h, desc,
                 size=13, color=SLATE, anchor=MSO_ANCHOR.MIDDLE)

    add_footer(s, 16)


def slide_thanks(prs):
    s = add_blank(prs)
    add_rect(s, 0, 0, SLIDE_W, SLIDE_H, NAVY)
    add_rect(s, 0, Inches(3.5), SLIDE_W, Inches(0.05), TEAL)
    add_text(s, Inches(0.8), Inches(2.4), Inches(11.7), Inches(0.8),
             "Thank you", size=44, bold=True, color=WHITE)
    add_text(s, Inches(0.8), Inches(3.7), Inches(11.7), Inches(0.5),
             "Harbourmaster  ·  Workflow Governance Control Plane",
             size=20, color=LIGHT)
    add_rounded(s, Inches(0.8), Inches(5.0), Inches(11.7), Inches(1.4), SLATE)
    add_text(s, Inches(1.0), Inches(5.15), Inches(11.3), Inches(1.2),
             "Try it locally:    make run    ·    make smoke    ·    make demo    ·    make ui\n"
             "Seed Elastic memory:    make index-elastic    ·    make redteam",
             size=14, color=WHITE, font="Consolas", anchor=MSO_ANCHOR.MIDDLE)


def build(path: Path) -> Path:
    prs = new_deck()
    slide_title(prs)
    slide_problem(prs)
    slide_solution(prs)
    slide_two_layers_table(prs)
    slide_architecture(prs)
    slide_langgraph_flow(prs)
    slide_specialists(prs)
    slide_verifier(prs)
    slide_human_loop(prs)
    slide_inline_guard(prs)
    slide_phoenix_elastic(prs)
    slide_policies(prs)
    slide_observability(prs)
    slide_deployment(prs)
    slide_key_features(prs)
    slide_stack(prs)
    slide_thanks(prs)
    prs.save(path)
    return path


if __name__ == "__main__":
    out = Path(__file__).resolve().parent.parent / "Harbourmaster_Overview.pptx"
    saved = build(out)
    print(f"Saved {saved}")
